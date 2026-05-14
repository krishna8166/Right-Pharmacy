"""
QuickMart — Web Application
===========================
Run locally:
    pip install flask
    python app.py
    → Customer : http://localhost:5000
    → Admin    : http://localhost:5000/admin   (password: right001)

Deploy free to Render / Railway:
    Add a Procfile:  web: gunicorn app:app
    pip install gunicorn, push repo, done.
"""

from flask import (Flask, render_template, request, jsonify,
                   session, redirect, url_for, send_file)
from functools import wraps
import sqlite3, json, io
from datetime import datetime
import os

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "right-pharmacy-secret-2026")

DB_PATH        = os.environ.get("DB_PATH", "retail_shop.db")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "right001")
SHOP_NAME      = "Right Pharmacy"

SAMPLE_ITEMS = [
    ("CTD-M 6.25/52",   120.00, 50),
    ("FEBUTEC-40",      210.00, 30),
    ("THYRONORM 50",    185.00, 100),
    ("NAPROSYN-D 500",  95.00,  40),
    ("CETISHAPE",       55.00,  60),
]

# ═══════════════════════════════════════════════════════
#  DATABASE
# ═══════════════════════════════════════════════════════
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as c:
        c.executescript("""
            CREATE TABLE IF NOT EXISTS items (
                id    INTEGER PRIMARY KEY AUTOINCREMENT,
                name  TEXT    NOT NULL,
                price REAL    NOT NULL,
                stock INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS orders (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                customer_name TEXT    NOT NULL,
                items_json    TEXT    NOT NULL,
                total         REAL    NOT NULL,
                status        TEXT    DEFAULT 'pending',
                created_at    TEXT    NOT NULL,
                delivered_at  TEXT
            );
            CREATE TABLE IF NOT EXISTS sales (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                order_id      INTEGER NOT NULL,
                customer_name TEXT    NOT NULL,
                items_json    TEXT    NOT NULL,
                amount        REAL    NOT NULL,
                sale_date     TEXT    NOT NULL,
                FOREIGN KEY(order_id) REFERENCES orders(id)
            );
        """)
        # ── Migrate: drop category column if it exists from a previous version ──
        cols = [r[1] for r in c.execute("PRAGMA table_info(items)").fetchall()]
        if "category" in cols:
            try:
                c.execute("ALTER TABLE items DROP COLUMN category")
            except Exception:
                c.executescript("""
                    CREATE TABLE items_new (
                        id    INTEGER PRIMARY KEY AUTOINCREMENT,
                        name  TEXT    NOT NULL,
                        price REAL    NOT NULL,
                        stock INTEGER NOT NULL DEFAULT 0
                    );
                    INSERT INTO items_new(id, name, price, stock)
                        SELECT id, name, price, stock FROM items;
                    DROP TABLE items;
                    ALTER TABLE items_new RENAME TO items;
                """)
            c.commit()
        # ── Seed only if table is empty ──────────────────────────────────────
        if c.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 0:
            c.executemany(
                "INSERT INTO items(name, price, stock) VALUES(?,?,?)",
                SAMPLE_ITEMS)
        c.commit()

# ═══════════════════════════════════════════════════════
#  AUTH
# ═══════════════════════════════════════════════════════
def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("admin_logged_in"):
            return redirect(url_for("admin_login_page"))
        return f(*args, **kwargs)
    return decorated

# ═══════════════════════════════════════════════════════
#  CUSTOMER ROUTES
# ═══════════════════════════════════════════════════════
@app.route("/")
def customer():
    return render_template("customer.html", shop_name=SHOP_NAME)

@app.route("/api/items")
def api_items():
    search = request.args.get("search", "").strip()
    q, p = "SELECT * FROM items", []
    if search:
        q += " WHERE name LIKE ?"
        p.append(f"%{search}%")
    q += " ORDER BY name"
    with get_db() as c:
        rows = c.execute(q, p).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/order", methods=["POST"])
def api_place_order():
    data  = request.get_json()
    name  = data.get("customer_name", "").strip()
    items = data.get("items", [])
    total = data.get("total", 0)
    if not name or not items:
        return jsonify({"error": "Missing name or items"}), 400
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db() as c:
        cur = c.execute(
            "INSERT INTO orders(customer_name,items_json,total,status,created_at)"
            " VALUES(?,?,?,?,?)",
            (name, json.dumps(items), total, "pending", now))
        c.commit()
    return jsonify({"order_id": cur.lastrowid})

# ═══════════════════════════════════════════════════════
#  ADMIN AUTH ROUTES
# ═══════════════════════════════════════════════════════
@app.route("/admin")
def admin_redirect():
    if session.get("admin_logged_in"):
        return redirect(url_for("admin_dashboard"))
    return redirect(url_for("admin_login_page"))

@app.route("/admin/login", methods=["GET"])
def admin_login_page():
    if session.get("admin_logged_in"):
        return redirect(url_for("admin_dashboard"))
    return render_template("admin.html", shop_name=SHOP_NAME, page="login")

@app.route("/admin/login", methods=["POST"])
def admin_login_post():
    data = request.get_json()
    if data and data.get("password") == ADMIN_PASSWORD:
        session["admin_logged_in"] = True
        return jsonify({"ok": True})
    return jsonify({"ok": False, "error": "Incorrect password"}), 401

@app.route("/admin/logout")
def admin_logout():
    session.pop("admin_logged_in", None)
    return redirect(url_for("admin_login_page"))

@app.route("/admin/dashboard")
@admin_required
def admin_dashboard():
    return render_template("admin.html", shop_name=SHOP_NAME, page="dashboard", now_date=datetime.now().strftime("%Y-%m-%d"))

# ═══════════════════════════════════════════════════════
#  ADMIN API ROUTES
# ═══════════════════════════════════════════════════════
@app.route("/api/admin/stats")
@admin_required
def api_stats():
    today = datetime.now().strftime("%Y-%m-%d")
    with get_db() as c:
        pending  = c.execute("SELECT COUNT(*) FROM orders WHERE status='pending'").fetchone()[0]
        revenue  = c.execute("SELECT COALESCE(SUM(amount),0) FROM sales WHERE sale_date=?",
                             (today,)).fetchone()[0]
        no_stock = c.execute("SELECT COUNT(*) FROM items WHERE stock=0").fetchone()[0]
    return jsonify({"pending": pending, "today_rev": revenue, "out_of_stock": no_stock})

@app.route("/api/admin/orders")
@admin_required
def api_orders():
    with get_db() as c:
        rows = c.execute(
            "SELECT * FROM orders WHERE status='pending'"
            " ORDER BY created_at DESC").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/admin/deliver/<int:oid>", methods=["POST"])
@admin_required
def api_deliver(oid):
    with get_db() as c:
        o = c.execute("SELECT * FROM orders WHERE id=?", (oid,)).fetchone()
        if not o:
            return jsonify({"error": "Order not found"}), 404
        if o["status"] != "pending":
            return jsonify({"error": "Already delivered"}), 400
        items = json.loads(o["items_json"])
        for it in items:
            c.execute("UPDATE items SET stock=MAX(0,stock-?) WHERE id=?",
                      (it["qty"], it["id"]))
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        c.execute("UPDATE orders SET status='delivered',delivered_at=? WHERE id=?",
                  (now, oid))
        c.execute(
            "INSERT INTO sales(order_id,customer_name,items_json,amount,sale_date)"
            " VALUES(?,?,?,?,?)",
            (oid, o["customer_name"], o["items_json"], o["total"], now[:10]))
        c.commit()
    return jsonify({"ok": True})

@app.route("/api/admin/inventory")
@admin_required
def api_inventory():
    search = request.args.get("search", "").strip()
    q, p = "SELECT * FROM items", []
    if search:
        q += " WHERE name LIKE ?"; p.append(f"%{search}%")
    q += " ORDER BY name"
    with get_db() as c:
        rows = c.execute(q, p).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/admin/stock/<int:iid>", methods=["POST"])
@admin_required
def api_stock(iid):
    data = request.get_json()
    with get_db() as c:
        if "value" in data:
            c.execute("UPDATE items SET stock=MAX(0,?) WHERE id=?",
                      (int(data["value"]), iid))
        elif "delta" in data:
            c.execute("UPDATE items SET stock=MAX(0,stock+?) WHERE id=?",
                      (int(data["delta"]), iid))
        c.commit()
        row = c.execute("SELECT stock FROM items WHERE id=?", (iid,)).fetchone()
    return jsonify({"ok": True, "stock": row["stock"] if row else 0})

@app.route("/api/admin/items/add", methods=["POST"])
@admin_required
def api_add_item():
    data  = request.get_json()
    name  = data.get("name", "").strip()
    price = data.get("price")
    stock = data.get("stock", 0)
    if not name or price is None:
        return jsonify({"error": "Name and price are required"}), 400
    try:
        price = float(price)
        stock = int(stock)
    except (ValueError, TypeError):
        return jsonify({"error": "Price and stock must be numbers"}), 400
    with get_db() as c:
        cur = c.execute(
            "INSERT INTO items(name, price, stock) VALUES(?,?,?)",
            (name, price, stock))
        c.commit()
        row = c.execute("SELECT * FROM items WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify({"ok": True, "item": dict(row)})

@app.route("/api/admin/items/delete/<int:iid>", methods=["POST"])
@admin_required
def api_delete_item(iid):
    with get_db() as c:
        row = c.execute("SELECT name FROM items WHERE id=?", (iid,)).fetchone()
        if not row:
            return jsonify({"error": "Item not found"}), 404
        c.execute("DELETE FROM items WHERE id=?", (iid,))
        c.commit()
    return jsonify({"ok": True})

@app.route("/api/admin/sales")
@admin_required
def api_sales():
    date = request.args.get("date", "").strip()
    q, p = "SELECT * FROM sales", []
    if date:
        q += " WHERE sale_date=?"; p.append(date)
    q += " ORDER BY id DESC"
    with get_db() as c:
        rows = c.execute(q, p).fetchall()
        total = c.execute(
            "SELECT COALESCE(SUM(amount),0) FROM sales" +
            (" WHERE sale_date=?" if date else ""),
            ([date] if date else [])).fetchone()[0]
    return jsonify({"sales": [dict(r) for r in rows], "total": total})

@app.route("/api/admin/items/template")
@admin_required
def api_download_template():
    """Return a pre-filled .xlsx template the admin can fill in."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return jsonify({"error": "openpyxl not installed"}), 500

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"

    # Header row styling
    hdr_fill = PatternFill("solid", fgColor="4361EE")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_border = Border(
        bottom=Side(style="medium", color="3047C8"))
    headers = ["Name", "Price (₹)", "Stock"]
    col_widths = [36, 14, 10]

    for col, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = hdr_border
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    # Sample rows so admin understands the format
    samples = [
        ("Crocin 500mg", 28.00, 100),
        ("Dolo 650",     32.00,  80),
        ("Paracetamol",  15.00, 150),
    ]
    note_font   = Font(italic=True, color="888888", size=10)
    data_border = Border(bottom=Side(style="thin", color="E2E8F0"))

    for r, (name, price, stock) in enumerate(samples, start=2):
        for col, val in enumerate([name, price, stock], start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.alignment = Alignment(horizontal="center" if col > 1 else "left",
                                       vertical="center")
            cell.border    = data_border
            if r == 2:                       # first data row hint
                cell.font = note_font

    # Freeze header row
    ws.freeze_panes = "A2"

    # Instructions sheet
    info = wb.create_sheet("Instructions")
    info["A1"] = "HOW TO USE THIS TEMPLATE"
    info["A1"].font = Font(bold=True, size=13)
    notes = [
        "",
        "1. Fill in the 'Items' sheet. Do not change the header row.",
        "2. Name    — required, any text (e.g. Crocin 500mg)",
        "3. Price   — required, number in ₹ (e.g. 28 or 28.50)",
        "4. Stock   — optional, defaults to 0 if left blank",
        "",
        "5. You may delete the sample rows before importing.",
        "6. Duplicate names will be SKIPPED (not overwritten).",
        "7. Rows with missing Name or invalid Price will be skipped",
        "   and listed in the error report after import.",
    ]
    for i, note in enumerate(notes, start=2):
        info.cell(row=i, column=1, value=note)
    info.column_dimensions["A"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name="import_template.xlsx")


@app.route("/api/admin/items/bulk-import", methods=["POST"])
@admin_required
def api_bulk_import():
    """Accept an .xlsx file and insert new items (skips duplicates by name)."""
    try:
        import openpyxl
    except ImportError:
        return jsonify({"error": "openpyxl is not installed on the server."}), 500

    file = request.files.get("file")
    if not file or not file.filename.lower().endswith((".xlsx", ".xls")):
        return jsonify({"error": "Please upload a valid .xlsx file."}), 400

    try:
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        return jsonify({"error": f"Could not read file: {e}"}), 400

    imported, skipped, errors = [], [], []

    # Detect header row — skip rows until we find Name / Price
    rows = list(ws.iter_rows(values_only=True))
    start = 0
    for i, row in enumerate(rows):
        vals = [str(v).strip().lower() if v else "" for v in row[:3]]
        if "name" in vals and any(k in vals for k in ("price", "price (₹)", "price (rs)")):
            start = i + 1   # data starts after header
            break

    with get_db() as c:
        existing_names = {
            r[0].strip().lower()
            for r in c.execute("SELECT name FROM items").fetchall()
        }

        for row_num, row in enumerate(rows[start:], start=start + 2):
            if not any(row):          # blank row — skip silently
                continue

            raw_name  = row[0] if len(row) > 0 else None
            raw_price = row[1] if len(row) > 1 else None
            raw_stock = row[2] if len(row) > 2 else 0

            name = str(raw_name).strip() if raw_name not in (None, "") else ""
            if not name:
                errors.append({"row": row_num, "reason": "Name is empty"})
                continue

            try:
                price = float(raw_price)
                if price < 0: raise ValueError()
            except (TypeError, ValueError):
                errors.append({"row": row_num, "name": name,
                               "reason": f"Invalid price '{raw_price}'"})
                continue

            try:
                stock = int(float(raw_stock)) if raw_stock not in (None, "") else 0
                stock = max(0, stock)
            except (TypeError, ValueError):
                stock = 0

            if name.lower() in existing_names:
                skipped.append({"row": row_num, "name": name,
                                "reason": "Name already exists"})
                continue

            c.execute("INSERT INTO items(name, price, stock) VALUES(?,?,?)",
                      (name, price, stock))
            existing_names.add(name.lower())
            imported.append({"name": name, "price": price, "stock": stock})

        c.commit()

    return jsonify({
        "ok":       True,
        "imported": len(imported),
        "skipped":  len(skipped),
        "errors":   len(errors),
        "items":    imported,
        "skipped_detail": skipped,
        "error_detail":   errors,
    })


# ═══════════════════════════════════════════════════════
#  STARTUP
# ═══════════════════════════════════════════════════════
init_db()
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"\n{'='*50}")
    print(f"  {SHOP_NAME} is running!")
    print(f"  Customer : http://localhost:{port}")
    print(f"  Admin    : http://localhost:{port}/admin")
    print(f"  Password : {ADMIN_PASSWORD}")
    print(f"{'='*50}\n")
    app.run(host="0.0.0.0", port=port, debug=False)
